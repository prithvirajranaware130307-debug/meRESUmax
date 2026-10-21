import io
import json
from reportlab.pdfgen import canvas
import docx
from app import create_app
from app.models import db, User, Resume, Analysis, Job

def run_verification():
    app = create_app()
    client = app.test_client()

    print('=== 1. HEALTH CHECK ===')
    res = client.get('/api/health')
    print('Health status:', res.status_code, res.get_json()['status'])
    assert res.status_code == 200

    print('\n=== 2. SEEDED DEMO CANDIDATE LOGIN ===')
    res = client.post('/api/auth/login', json={'email': 'demo.candidate@meresumax.ai', 'password': 'DemoPass123!'})
    print('Demo login status:', res.status_code)
    assert res.status_code == 200
    demo_token = res.get_json()['token']
    auth_headers = {'Authorization': f'Bearer {demo_token}'}

    print('\n=== 3. USER REGISTRATION (NEW ACCOUNT) ===')
    email = 'demo_verifier@meresumax.ai'
    reg_res = client.post('/api/auth/register', json={
        'name': 'Alex Demo',
        'email': email,
        'password': 'SecurePassword123!',
        'target_role': 'Python Developer'
    })
    print('Registration status:', reg_res.status_code)
    assert reg_res.status_code in (201, 409)
    if reg_res.status_code == 201:
        token = reg_res.get_json()['token']
    else:
        login_res = client.post('/api/auth/login', json={'email': email, 'password': 'SecurePassword123!'})
        token = login_res.get_json()['token']
    headers = {'Authorization': f'Bearer {token}'}

    print('\n=== 4. GET /api/auth/me (PROTECTED ROUTE & SESSION) ===')
    me_res = client.get('/api/auth/me', headers=headers)
    print('Me status:', me_res.status_code, me_res.get_json()['user']['email'])
    assert me_res.status_code == 200

    print('\n=== 5A. RESUME UPLOAD (PDF FILE) ===')
    pdf_buffer = io.BytesIO()
    p = canvas.Canvas(pdf_buffer)
    p.drawString(100, 750, 'Alex Rivera')
    p.drawString(100, 730, 'Python Developer - alex@rivera.dev - (555) 019-2834')
    p.drawString(100, 710, 'Skills: Python, Flask, Django, PostgreSQL, Docker, AWS, Redis, Git, Unit Testing, REST APIs')
    p.drawString(100, 690, 'Experience: Backend Software Engineer at Nexus Cloud Solutions (2021 - Present)')
    p.drawString(100, 670, '- Architected 15+ REST microservices in Python Flask, reducing API latency by 35%.')
    p.drawString(100, 650, '- Designed and indexed PostgreSQL schemas handling over 2M records.')
    p.drawString(100, 630, '- Built Dockerized CI/CD deployment pipelines on AWS ECS.')
    p.drawString(100, 610, 'Education: B.S. in Computer Science, State University, 2021')
    p.save()
    pdf_buffer.seek(0)

    upload_pdf = client.post(
        '/api/resumes/upload',
        data={'file': (pdf_buffer, 'alex_rivera_resume.pdf')},
        headers=headers,
        content_type='multipart/form-data'
    )
    print('Upload PDF status:', upload_pdf.status_code)
    assert upload_pdf.status_code == 201
    resume_id = upload_pdf.get_json()['resume']['id']
    print('Resume uploaded successfully with ID:', resume_id)

    print('\n=== 5B. RESUME UPLOAD (DOCX FILE) ===')
    doc = docx.Document()
    doc.add_heading('Taylor Swift - Frontend Engineer', 0)
    doc.add_paragraph('Skills: React, TypeScript, JavaScript, CSS, HTML, Tailwind, Next.js, Redux, Jest')
    doc.add_paragraph('Experience: Built high-performance responsive web applications serving 100k users.')
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    upload_docx = client.post(
        '/api/resumes/upload',
        data={'file': (docx_buffer, 'taylor_resume.docx')},
        headers=headers,
        content_type='multipart/form-data'
    )
    print('Upload DOCX status:', upload_docx.status_code)
    assert upload_docx.status_code == 201
    print('DOCX parsed successfully, word count:', upload_docx.get_json()['resume']['parsed_data']['word_count'])

    print('\n=== 5C. RESUME ANALYSIS GENERATION ===')
    analysis_res = client.post(
        '/api/analyses',
        json={'resume_id': resume_id, 'target_role': 'Python Developer'},
        headers=headers
    )
    print('Create Analysis status:', analysis_res.status_code)
    assert analysis_res.status_code == 201
    analysis_data = analysis_res.get_json()['analysis']
    print('Quality Score:', analysis_data['quality_score'])
    print('Job Match Score:', analysis_data['job_match_score'])
    print('Category Scores:', analysis_data['category_scores'])
    print('Matched Skills:', analysis_data['matched_skills'])
    print('Missing Skills:', analysis_data['missing_skills'])
    print('AI Feedback is_live_ai:', analysis_data['ai_feedback'].get('is_live_ai'))
    assert analysis_data['quality_score'] > 0

    print('\n=== 6. JOBS RECOMMENDATIONS (ROLE-FILTERED & CURATED FALLBACK) ===')
    jobs_res = client.get('/api/jobs?role=Python%20Developer', headers=headers)
    print('Jobs status:', jobs_res.status_code)
    assert jobs_res.status_code == 200
    jobs = jobs_res.get_json()['jobs']
    print('Number of jobs returned:', len(jobs))
    print('Top Job:', jobs[0]['title'], '| Company:', jobs[0]['company'], '| Location:', jobs[0]['location'])
    assert len(jobs) >= 5

    print('\n=== 7. AI INTERVIEW QUESTIONS GENERATION (ALL 5 REQUIRED ROLES) ===')
    test_roles = ['Frontend Developer', 'Backend Developer', 'Python Developer', 'Data Analyst', 'AI/ML Engineer']
    for test_role in test_roles:
        q_res = client.post('/api/interview/questions', json={'target_role': test_role, 'difficulty': 'Mid-Level'}, headers=headers)
        assert q_res.status_code == 200
        qs = q_res.get_json()['questions']
        q_text = qs[0]['question'][:50]
        print(f"Role [{test_role}]: generated {len(qs)} questions. Sample: {q_text}...")
        assert len(qs) == 5

    print('\n=== 8. AI INTERVIEW ANSWER EVALUATION ===')
    eval_res = client.post('/api/interview/evaluate', json={
        'question': 'Explain Python Global Interpreter Lock and how to achieve concurrency.',
        'category': 'Technical Depth',
        'answer': 'The GIL prevents multiple native threads from executing Python bytecodes concurrently. For CPU-bound tasks, we use the multiprocessing module or ProcessPoolExecutor to run tasks across separate cores with distinct memory spaces. For I/O-bound tasks, asyncio with non-blocking event loops or thread pools allow high concurrent connection throughput without GIL contention.',
        'target_role': 'Python Developer',
        'difficulty': 'Mid-Level'
    }, headers=headers)
    print('Evaluation status:', eval_res.status_code)
    assert eval_res.status_code == 200
    eval_json = eval_res.get_json()
    evaluation = eval_json.get('evaluation', eval_json)
    print('Overall Score:', evaluation['overall_score'])
    print('STAR Score:', evaluation['star_score'])
    print('Strengths count:', len(evaluation['strengths']))
    print('Improvements count:', len(evaluation['areas_for_improvement']))
    better_answer_snippet = evaluation.get('suggested_exemplar_answer', '')[:70]
    print(f"Suggested rewrite snippet: {better_answer_snippet}...")
    assert evaluation['overall_score'] >= 40

    print('\n=== 9. FLOATING ASSISTANT CHATBOT ===')
    chat_res = client.post('/api/assistant/chat', json={
        'message': 'Can you help me improve my resume for Python Developer roles?',
        'context': {'target_role': 'Python Developer', 'score': analysis_data['quality_score']}
    }, headers=headers)
    print('Chatbot status:', chat_res.status_code)
    assert chat_res.status_code == 200
    chat_json = chat_res.get_json()
    print('Chatbot Provider:', chat_json['provider'])
    reply_snippet = chat_json['reply'][:100]
    print(f"Chatbot Reply: {reply_snippet}...")
    assert len(chat_json['reply']) > 20

    print('\n=== 10. AUTH LOGOUT ===')
    logout_res = client.post('/api/auth/logout', headers=headers)
    print('Logout status:', logout_res.status_code)
    assert logout_res.status_code == 200

    print('\n' + '='*60)
    print('>>> SUCCESS: ALL REQUIRED FLOWS VERIFIED LOCALLY! <<<')
    print('='*60)

if __name__ == '__main__':
    run_verification()
